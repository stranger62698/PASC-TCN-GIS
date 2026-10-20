from __future__ import annotations

import re
import sys
from pathlib import Path

from PIL import Image
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "lanjifyw-product-case-study.md"
OUTPUT = ROOT / "docs" / "lanjifyw-product-case-study.docx"
QA_DIR = ROOT / "docs" / ".qa" / "lanjifyw-product-case-study"

NAVY = "0B2545"
BLUE = "1677FF"
MID_BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
MUTED = "5F708A"
PALE_BLUE = "EAF3FF"
TABLE_FILL = "F4F6F9"
WHITE = "FFFFFF"
BODY_FONT = "PingFang SC"
LATIN_FONT = "Inter"

# narrative_proposal preset + named LANJIFYW brand overrides.
TOKENS = {
    "page": {"width": 8.5, "height": 11.0, "margin": 1.0, "header": 0.492, "footer": 0.492},
    "body": {"size": 10.5, "after": 8, "line": 1.333},
    "title": {"size": 29, "color": NAVY, "after": 8},
    "h1": {"size": 16, "color": BLUE, "before": 18, "after": 10},
    "h2": {"size": 13, "color": MID_BLUE, "before": 12, "after": 6},
    "h3": {"size": 12, "color": DARK_BLUE, "before": 8, "after": 4},
    "list": {"left": 0.375, "hanging": 0.194, "after": 4, "line": 1.208},
    "table": {"width_dxa": 9360, "indent_dxa": 120, "cell_start_dxa": 120},
}


def rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value)


def set_run_font(run, name=BODY_FONT, size=None, color=NAVY, bold=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), LATIN_FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), LATIN_FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = rgb(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa):
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(TOKENS["table"]["indent_dxa"]))
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            width = widths_dxa[index]
            tc_w = cell._tc.get_or_add_tcPr().find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                cell._tc.get_or_add_tcPr().append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_field(paragraph, instruction):
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run = paragraph.add_run()
    run._r.extend([begin, instr, separate, text, end])
    set_run_font(run, size=9, color=MUTED)


def configure_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = Pt(TOKENS["body"]["size"])
    normal.font.color.rgb = rgb(NAVY)
    normal._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), BODY_FONT)
    normal._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), LATIN_FONT)
    normal._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), LATIN_FONT)
    pf = normal.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(TOKENS["body"]["after"])
    pf.line_spacing = TOKENS["body"]["line"]
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    for name, token in (("Heading 1", TOKENS["h1"]), ("Heading 2", TOKENS["h2"]), ("Heading 3", TOKENS["h3"])):
        style = styles[name]
        style.font.name = BODY_FONT
        style.font.size = Pt(token["size"])
        style.font.bold = True
        style.font.color.rgb = rgb(token["color"])
        style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), BODY_FONT)
        style._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), LATIN_FONT)
        style._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), LATIN_FONT)
        style.paragraph_format.space_before = Pt(token["before"])
        style.paragraph_format.space_after = Pt(token["after"])
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True

    for name in ("List Bullet", "List Number"):
        style = styles[name]
        style.font.name = BODY_FONT
        style.font.size = Pt(TOKENS["body"]["size"])
        style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), BODY_FONT)
        style._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), LATIN_FONT)
        style._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), LATIN_FONT)
        style.paragraph_format.left_indent = Inches(TOKENS["list"]["left"])
        style.paragraph_format.first_line_indent = Inches(-TOKENS["list"]["hanging"])
        style.paragraph_format.space_after = Pt(TOKENS["list"]["after"])
        style.paragraph_format.line_spacing = TOKENS["list"]["line"]

    caption = styles["Caption"]
    caption.font.name = BODY_FONT
    caption.font.size = Pt(9)
    caption.font.italic = False
    caption.font.color.rgb = rgb(MUTED)
    caption._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), BODY_FONT)
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_before = Pt(3)
    caption.paragraph_format.space_after = Pt(10)
    caption.paragraph_format.keep_with_next = True

    if "Lead Callout" not in styles:
        callout = styles.add_style("Lead Callout", WD_STYLE_TYPE.PARAGRAPH)
    else:
        callout = styles["Lead Callout"]
    callout.base_style = normal
    callout.font.name = BODY_FONT
    callout.font.size = Pt(11)
    callout.font.bold = True
    callout.font.color.rgb = rgb(DARK_BLUE)
    callout.paragraph_format.left_indent = Inches(0.18)
    callout.paragraph_format.right_indent = Inches(0.18)
    callout.paragraph_format.space_before = Pt(8)
    callout.paragraph_format.space_after = Pt(12)
    callout.paragraph_format.line_spacing = 1.25


def configure_page(doc):
    section = doc.sections[0]
    section.page_width = Inches(TOKENS["page"]["width"])
    section.page_height = Inches(TOKENS["page"]["height"])
    section.top_margin = Inches(TOKENS["page"]["margin"])
    section.right_margin = Inches(TOKENS["page"]["margin"])
    section.bottom_margin = Inches(TOKENS["page"]["margin"])
    section.left_margin = Inches(TOKENS["page"]["margin"])
    section.header_distance = Inches(TOKENS["page"]["header"])
    section.footer_distance = Inches(TOKENS["page"]["footer"])
    section.different_first_page_header_footer = True

    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run("澜迹 LANJIFYW  /  PRODUCT CASE STUDY")
    set_run_font(r, LATIN_FONT, 8.5, MUTED, bold=True)

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    fp.paragraph_format.space_before = Pt(0)
    fr = fp.add_run("产品案例说明    ")
    set_run_font(fr, BODY_FONT, 8.5, MUTED)
    add_field(fp, "PAGE")


def add_cover(doc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(13)
    r = p.add_run("PRODUCT CASE STUDY")
    set_run_font(r, LATIN_FONT, 10, BLUE, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(TOKENS["title"]["after"])
    r = p.add_run("澜迹 LANJIFYW")
    set_run_font(r, BODY_FONT, TOKENS["title"]["size"], TOKENS["title"]["color"], bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(18)
    r = p.add_run("时序 InSAR 可解释分析与现场证据联动平台")
    set_run_font(r, BODY_FONT, 15, DARK_BLUE, bold=True)

    image_path = ROOT / "public" / "hero-insar-observation-v2.webp"
    converted = QA_DIR / "hero-cover.png"
    converted.parent.mkdir(parents=True, exist_ok=True)
    with Image.open(image_path) as im:
        im.convert("RGB").save(converted, "PNG")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(12)
    run = p.add_run()
    run.add_picture(str(converted), width=Inches(5.9))
    set_picture_alt(run, "SAR 卫星形变观测与时空分析链路主视觉")

    p = doc.add_paragraph(style="Lead Callout")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("把数万监测点与上百期观测，转化为可定位、可比较、可解释、可继续核查的证据。")
    set_run_font(r, BODY_FONT, 11, DARK_BLUE, bold=True)
    shade_paragraph(p, PALE_BLUE)

    for label, value in (
        ("定位", "时空地图 + 形变模式 + LBS 现场证据"),
        ("案例", "城市建筑 / 道路 / 滑坡"),
        ("交付", "真实数据案例、分析工作台与可复核边界"),
    ):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(f"{label}  ")
        set_run_font(r, BODY_FONT, 9.5, BLUE, bold=True)
        r = p.add_run(value)
        set_run_font(r, BODY_FONT, 9.5, MUTED)

    doc.add_page_break()


def shade_paragraph(paragraph, fill):
    p_pr = paragraph._p.get_or_add_pPr()
    shd = p_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        p_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_picture_alt(run, description):
    drawings = run._r.findall(qn("w:drawing"))
    for drawing in drawings:
        for doc_pr in drawing.iter(qn("wp:docPr")):
            doc_pr.set("descr", description)


def add_inline_markup(paragraph, text):
    parts = re.split(r"(\*\*.*?\*\*|\*.*?\*)", text)
    for part in parts:
        if not part:
            continue
        bold = part.startswith("**") and part.endswith("**")
        italic = not bold and part.startswith("*") and part.endswith("*")
        content = part[2:-2] if bold else part[1:-1] if italic else part
        run = paragraph.add_run(content)
        set_run_font(run, BODY_FONT, TOKENS["body"]["size"], NAVY, bold=bold, italic=italic)


def resolve_image(path_text):
    path = (SOURCE.parent / path_text).resolve()
    if path.suffix.lower() == ".webp":
        converted = QA_DIR / f"{path.stem}.png"
        converted.parent.mkdir(parents=True, exist_ok=True)
        with Image.open(path) as im:
            im.convert("RGB").save(converted, "PNG")
        return converted
    return path


def add_image(doc, path_text, alt):
    path = resolve_image(path_text)
    with Image.open(path) as im:
        width, height = im.size
    max_w = 5.9
    max_h = 4.3
    ratio = min(max_w / width, max_h / height) * 96
    width_in = min(max_w, width * ratio / 96)
    height_in = min(max_h, height * ratio / 96)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    run = p.add_run()
    run.add_picture(str(path), width=Inches(width_in), height=Inches(height_in))
    set_picture_alt(run, alt)


def add_markdown_table(doc, rows):
    headers = rows[0]
    body = rows[2:]
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    for index, value in enumerate(headers):
        cell = table.rows[0].cells[index]
        set_cell_shading(cell, TABLE_FILL)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(value)
        set_run_font(r, BODY_FONT, 9, NAVY, bold=True)
    for source_row in body:
        row = table.add_row()
        for index, value in enumerate(source_row):
            cell = row.cells[index]
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if index in (0, len(source_row) - 1) else WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.15
            r = p.add_run(value)
            set_run_font(r, BODY_FONT, 8.7, NAVY)
    widths = [1500, 2620, 2620, 2620] if len(headers) == 4 else [9360 // len(headers)] * len(headers)
    widths[-1] += 9360 - sum(widths)
    set_table_geometry(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def create_restarted_decimal_numbering(doc):
    root = doc.part.numbering_part.element
    abstract_id = None
    for abstract in root.findall(qn("w:abstractNum")):
        num_fmt = abstract.find(".//" + qn("w:numFmt"))
        if num_fmt is not None and num_fmt.get(qn("w:val")) == "decimal":
            abstract_id = abstract.get(qn("w:abstractNumId"))
            break
    if abstract_id is None:
        raise RuntimeError("No decimal numbering definition is available.")
    existing = [int(node.get(qn("w:numId"))) for node in root.findall(qn("w:num"))]
    num_id = max(existing, default=0) + 1
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), abstract_id)
    num.append(abstract_ref)
    override = OxmlElement("w:lvlOverride")
    override.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:startOverride")
    start.set(qn("w:val"), "1")
    override.append(start)
    num.append(override)
    root.append(num)
    return num_id


def apply_numbering(paragraph, num_id):
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = p_pr.find(qn("w:numPr"))
    if num_pr is None:
        num_pr = OxmlElement("w:numPr")
        p_pr.append(num_pr)
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num = OxmlElement("w:numId")
    num.set(qn("w:val"), str(num_id))
    num_pr.extend([ilvl, num])


def parse_table(lines, start):
    rows = []
    index = start
    while index < len(lines) and lines[index].strip().startswith("|"):
        rows.append([cell.strip() for cell in lines[index].strip().strip("|").split("|")])
        index += 1
    return rows, index


def add_body_from_markdown(doc):
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    index = next(i for i, line in enumerate(lines) if line.strip() == "---") + 1
    active_num_id = None
    while index < len(lines):
        raw = lines[index]
        line = raw.strip()
        # Named override: closing-note de-duplication. The provenance/safety caveat
        # is already stated in the case boundaries; omitting this repeated Markdown
        # note from Word avoids an otherwise orphaned final page.
        if line.startswith("**说明：** 数据来自项目资产与回归验证"):
            active_num_id = None
            index += 1
            continue
        if not line or line == "---":
            active_num_id = None
            index += 1
            continue
        if line.startswith("!["):
            active_num_id = None
            match = re.match(r"!\[(.*?)\]\((.*?)\)", line)
            if match:
                add_image(doc, match.group(2), match.group(1))
            index += 1
            continue
        if line.startswith("|"):
            active_num_id = None
            rows, index = parse_table(lines, index)
            add_markdown_table(doc, rows)
            continue
        if line.startswith("### "):
            active_num_id = None
            p = doc.add_paragraph(line[4:], style="Heading 2")
        elif line.startswith("## "):
            active_num_id = None
            p = doc.add_paragraph(line[3:], style="Heading 1")
        elif line.startswith("# "):
            active_num_id = None
            p = doc.add_paragraph(line[2:], style="Heading 1")
        elif re.match(r"^- ", line):
            active_num_id = None
            p = doc.add_paragraph(style="List Bullet")
            add_inline_markup(p, line[2:])
        elif re.match(r"^\d+\. ", line):
            if active_num_id is None:
                active_num_id = create_restarted_decimal_numbering(doc)
            p = doc.add_paragraph(style="List Number")
            apply_numbering(p, active_num_id)
            add_inline_markup(p, re.sub(r"^\d+\. ", "", line))
        elif line.startswith(">"):
            active_num_id = None
            p = doc.add_paragraph(style="Lead Callout")
            add_inline_markup(p, line.lstrip("> "))
            shade_paragraph(p, PALE_BLUE)
        elif line.startswith("*图 ") and line.endswith("*"):
            active_num_id = None
            p = doc.add_paragraph(line[1:-1], style="Caption")
        else:
            active_num_id = None
            p = doc.add_paragraph()
            add_inline_markup(p, line)
        index += 1


def audit_document(doc):
    section = doc.sections[0]
    assert round(section.page_width.inches, 2) == 8.5
    assert round(section.left_margin.inches, 2) == 1.0
    assert round(section.right_margin.inches, 2) == 1.0
    assert doc.styles["Normal"].font.size.pt == TOKENS["body"]["size"]
    assert doc.styles["Heading 1"].font.size.pt == TOKENS["h1"]["size"]
    for table in doc.tables:
        tbl_pr = table._tbl.tblPr
        assert tbl_pr.find(qn("w:tblW")).get(qn("w:type")) == "dxa"
        assert tbl_pr.find(qn("w:tblInd")).get(qn("w:w")) == str(TOKENS["table"]["indent_dxa"])


def build():
    QA_DIR.mkdir(parents=True, exist_ok=True)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    doc.core_properties.title = "澜迹 LANJIFYW - 时序 InSAR 产品案例说明"
    doc.core_properties.subject = "产品案例、作品集与产品经理简历素材"
    doc.core_properties.author = "LANJIFYW"
    doc.core_properties.keywords = "InSAR, WebGIS, 时空地图, 形变模式, LBS, 产品案例"
    configure_styles(doc)
    configure_page(doc)
    add_cover(doc)
    add_body_from_markdown(doc)
    audit_document(doc)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    try:
        build()
    except Exception as exc:
        print(f"build failed: {exc}", file=sys.stderr)
        raise
